"""Builds ml/pallor/data/prepared/{anemic,non_anemic}/ from the Eyes-defy-anemia
dataset's raw layout (India/Italy patient folders + India.xlsx/Italy.xlsx
lab values), ready for shared.coreml_convert.load_image_dataset.

Label: WHO anemia thresholds by gender (Hgb < 13 g/dL for men, < 12 g/dL
for women) applied to each patient's lab hemoglobin value — not a
Kaggle-provided binary label, so this is real clinical ground truth rather
than a proxy.

Image: each patient's "<id>_palpebral.png" — a pre-segmented crop of the
palpebral conjunctiva (the standard clinical site for anemia screening: the
inside of the pulled-down lower eyelid), already isolated by the dataset's
own creators. Excludes "*_forniceal_palpebral.png" (a different crop some
India patients also have).
"""
import shutil
from pathlib import Path

import openpyxl

RAW_DIR = Path(__file__).parent / "data" / "raw" / "eyes_defy_anemia" / "dataset anemia"
OUT_DIR = Path(__file__).parent / "data" / "prepared"

COHORTS = ["India", "Italy"]


def load_labels(cohort: str) -> dict:
    wb = openpyxl.load_workbook(RAW_DIR / cohort / f"{cohort}.xlsx")
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip().lower() if h else "" for h in rows[0]]
    number_idx, hgb_idx, gender_idx = header.index("number"), header.index("hgb"), header.index("gender")

    labels = {}
    for row in rows[1:]:
        if row[number_idx] is None or row[hgb_idx] is None or row[gender_idx] is None:
            continue
        patient_id = str(int(row[number_idx]))
        try:
            # A few Italian-cohort rows use a comma decimal separator (e.g. "15,1");
            # a few cells are placeholders like "_" for missing/unusable lab values.
            hgb = float(str(row[hgb_idx]).replace(",", "."))
        except ValueError:
            continue
        gender = str(row[gender_idx]).strip().upper()
        threshold = 13.0 if gender == "M" else 12.0
        labels[patient_id] = "anemic" if hgb < threshold else "non_anemic"
    return labels


def find_palpebral_image(patient_dir: Path) -> Path | None:
    candidates = [p for p in patient_dir.glob("*_palpebral.png") if "forniceal" not in p.name]
    return candidates[0] if candidates else None


def main() -> None:
    for label in ["anemic", "non_anemic"]:
        (OUT_DIR / label).mkdir(parents=True, exist_ok=True)

    counts = {"anemic": 0, "non_anemic": 0, "missing_image": 0, "missing_label": 0}
    for cohort in COHORTS:
        labels = load_labels(cohort)
        cohort_dir = RAW_DIR / cohort
        for patient_dir in sorted(p for p in cohort_dir.iterdir() if p.is_dir()):
            patient_id = patient_dir.name
            label = labels.get(patient_id)
            if label is None:
                counts["missing_label"] += 1
                continue
            image = find_palpebral_image(patient_dir)
            if image is None:
                counts["missing_image"] += 1
                continue
            dest = OUT_DIR / label / f"{cohort.lower()}_{patient_id}{image.suffix}"
            shutil.copy2(image, dest)
            counts[label] += 1

    print(counts)


if __name__ == "__main__":
    main()
